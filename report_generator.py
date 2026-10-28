"""Professional DOCX Comparison Report Generator.

Transforms validated ComparisonResult data into a structured, executive-ready
Microsoft Word (.docx) document with custom styling, comparison tables,
evidence citations, and disclaimers.
"""

from datetime import datetime
import io
from typing import List, Optional, Union
import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from config import APP_DISCLAIMER, APP_TITLE
from schemas import ComparisonResult


def set_cell_background(cell, hex_color: str):
    """Sets background shading color for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding/margins for a table cell in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for margin_name, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{margin_name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def format_page_list(pages: List[int]) -> str:
    """Formats list of page integers into readable string."""
    if not pages:
        return "Not specified"
    if len(pages) == 1:
        return f"p. {pages[0]}"
    return f"pp. {', '.join(str(p) for p in pages)}"


def generate_docx_report(
    result: ComparisonResult,
    output_path: Optional[str] = None,
) -> io.BytesIO:
    """Generates a styled DOCX report from ComparisonResult.
    
    Args:
        result: Validated ComparisonResult schema.
        output_path: Optional file path to persist the generated document.
        
    Returns:
        BytesIO buffer containing the DOCX file data.
    """
    doc = Document()

    # Define Palette
    COLOR_PRIMARY = RGBColor(27, 54, 93)     # Deep Navy #1B365D
    COLOR_SECONDARY = RGBColor(74, 85, 104)  # Slate Gray #4A5568
    COLOR_MUTED = RGBColor(113, 128, 150)    # Cool Gray #718096
    COLOR_TEXT = RGBColor(45, 55, 72)        # Charcoal #2D3748
    HEX_HEADER_BG = "1B365D"
    HEX_ALT_ROW_BG = "F7FAFC"
    HEX_CALLOUT_BG = "EDF2F7"

    # Set document standard margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Base Normal Style
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_TEXT
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # --- COVER / HEADER SECTION ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(2)
    title_p.paragraph_format.space_before = Pt(0)
    title_run = title_p.add_run(APP_TITLE)
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    sub_run = sub_p.add_run("Executive Comparative Intelligence & Audit Report")
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = COLOR_SECONDARY
    sub_run.italic = True

    # Metadata Summary Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_items = [
        ("Document A (Baseline):", result.doc_a_name or "Document A"),
        ("Document B (Comparison):", result.doc_b_name or "Document B"),
        ("Generated Timestamp:", result.timestamp or datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("AI Model Engine:", result.model_used or "Gemini API"),
    ]

    for idx, (label, val) in enumerate(meta_items):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.8)

        set_cell_background(cell_lbl, "F0F4F8")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(0)
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.color.rgb = COLOR_PRIMARY

        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(0)
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)
        r_val.font.color.rgb = COLOR_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_section_header(title: str, level: int = 1):
        """Helper to create branded section titles."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title)
        run.bold = True
        if level == 1:
            run.font.size = Pt(14)
            run.font.color.rgb = COLOR_PRIMARY
        else:
            run.font.size = Pt(12)
            run.font.color.rgb = COLOR_SECONDARY
        return p

    # --- 1. EXECUTIVE SUMMARIES ---
    add_section_header("1. Executive Document Summaries")

    # Document A Summary Box
    doc_a_label = doc.add_paragraph()
    doc_a_label.paragraph_format.space_after = Pt(2)
    doc_a_r = doc_a_label.add_run(f"Summary: Document A ({result.doc_a_name or 'Document A'})")
    doc_a_r.bold = True
    doc_a_r.font.size = Pt(11)
    doc_a_r.font.color.rgb = COLOR_PRIMARY

    p_a = doc.add_paragraph()
    p_a.paragraph_format.space_after = Pt(10)
    p_a.paragraph_format.left_indent = Inches(0.2)
    p_a.add_run(result.document_a_summary or "No summary available for Document A.")

    # Document B Summary Box
    doc_b_label = doc.add_paragraph()
    doc_b_label.paragraph_format.space_after = Pt(2)
    doc_b_r = doc_b_label.add_run(f"Summary: Document B ({result.doc_b_name or 'Document B'})")
    doc_b_r.bold = True
    doc_b_r.font.size = Pt(11)
    doc_b_r.font.color.rgb = COLOR_PRIMARY

    p_b = doc.add_paragraph()
    p_b.paragraph_format.space_after = Pt(14)
    p_b.paragraph_format.left_indent = Inches(0.2)
    p_b.add_run(result.document_b_summary or "No summary available for Document B.")

    # --- 2. KEY SIMILARITIES ---
    add_section_header("2. Core Similarities & Shared Provisions")
    if not result.similarities:
        doc.add_paragraph("No significant common provisions or identical clauses were detected.")
    else:
        for idx, sim in enumerate(result.similarities, start=1):
            p_topic = doc.add_paragraph()
            p_topic.paragraph_format.space_before = Pt(6)
            p_topic.paragraph_format.space_after = Pt(2)
            r_num = p_topic.add_run(f"{idx}. {sim.topic}")
            r_num.bold = True
            r_num.font.size = Pt(11)
            r_num.font.color.rgb = COLOR_PRIMARY

            p_desc = doc.add_paragraph()
            p_desc.paragraph_format.left_indent = Inches(0.25)
            p_desc.paragraph_format.space_after = Pt(4)
            p_desc.add_run(sim.description)

            if sim.evidence_a or sim.evidence_b:
                p_ev = doc.add_paragraph()
                p_ev.paragraph_format.left_indent = Inches(0.4)
                p_ev.paragraph_format.space_after = Pt(6)
                p_ev_run = p_ev.add_run(
                    f"Evidence A ({format_page_list(sim.pages_a)}): \"{sim.evidence_a}\"\n"
                    f"Evidence B ({format_page_list(sim.pages_b)}): \"{sim.evidence_b}\""
                )
                p_ev_run.font.size = Pt(9.5)
                p_ev_run.font.color.rgb = COLOR_MUTED
                p_ev_run.italic = True

    # --- 3. MEANINGFUL DIFFERENCES ---
    add_section_header("3. Meaningful Differences & Modifications")
    if not result.differences:
        doc.add_paragraph("No substantive differences or conflicts were detected between the documents.")
    else:
        diff_table = doc.add_table(rows=1, cols=4)
        diff_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        diff_table.autofit = False

        headers = ["Topic & Change Type", "Document A Statement", "Document B Statement", "Significance & Impact"]
        widths = [Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.6)]

        hdr_row = diff_table.rows[0]
        for idx, text in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.width = widths[idx]
            set_cell_background(cell, HEX_HEADER_BG)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, diff in enumerate(result.differences):
            row = diff_table.add_row()
            bg_color = HEX_ALT_ROW_BG if row_idx % 2 == 1 else "FFFFFF"

            # Column 1: Topic + Change Type
            c0 = row.cells[0]
            c0.width = widths[0]
            set_cell_background(c0, bg_color)
            set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(2)
            r0_topic = p0.add_run(f"{diff.topic}\n")
            r0_topic.bold = True
            r0_topic.font.size = Pt(9.5)
            r0_badge = p0.add_run(f"[{diff.change_type.upper()}]")
            r0_badge.bold = True
            r0_badge.font.size = Pt(8.5)
            r0_badge.font.color.rgb = COLOR_SECONDARY

            # Column 2: Document A
            c1 = row.cells[1]
            c1.width = widths[1]
            set_cell_background(c1, bg_color)
            set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(diff.document_a_value or "Not stated")
            r1.font.size = Pt(9)
            if diff.evidence_a:
                p1_ev = c1.add_paragraph()
                p1_ev.paragraph_format.space_after = Pt(0)
                r1_ev = p1_ev.add_run(f"Cite ({format_page_list(diff.pages_a)}): \"{diff.evidence_a}\"")
                r1_ev.font.size = Pt(8)
                r1_ev.font.color.rgb = COLOR_MUTED
                r1_ev.italic = True

            # Column 3: Document B
            c2 = row.cells[2]
            c2.width = widths[2]
            set_cell_background(c2, bg_color)
            set_cell_margins(c2, top=80, bottom=80, left=120, right=120)
            p2 = c2.paragraphs[0]
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(diff.document_b_value or "Not stated")
            r2.font.size = Pt(9)
            if diff.evidence_b:
                p2_ev = c2.add_paragraph()
                p2_ev.paragraph_format.space_after = Pt(0)
                r2_ev = p2_ev.add_run(f"Cite ({format_page_list(diff.pages_b)}): \"{diff.evidence_b}\"")
                r2_ev.font.size = Pt(8)
                r2_ev.font.color.rgb = COLOR_MUTED
                r2_ev.italic = True

            # Column 4: Impact
            c3 = row.cells[3]
            c3.width = widths[3]
            set_cell_background(c3, bg_color)
            set_cell_margins(c3, top=80, bottom=80, left=120, right=120)
            p3 = c3.paragraphs[0]
            p3.paragraph_format.space_after = Pt(0)
            r3 = p3.add_run(diff.impact or "N/A")
            r3.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- 4. KEY DETAILS COMPARISON ---
    add_section_header("4. Key Details Comparison (Dates, Fees, Requirements)")
    if not result.key_details:
        doc.add_paragraph("No specific factual criteria or metrics were matched for side-by-side comparison.")
    else:
        kd_table = doc.add_table(rows=1, cols=4)
        kd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        kd_table.autofit = False

        kd_headers = ["Attribute / Field", "Document A Baseline", "Document B Comparison", "Status"]
        kd_widths = [Inches(2.0), Inches(2.2), Inches(2.2), Inches(1.0)]

        kd_hdr_row = kd_table.rows[0]
        for idx, text in enumerate(kd_headers):
            cell = kd_hdr_row.cells[idx]
            cell.width = kd_widths[idx]
            set_cell_background(cell, "2D3748")  # Dark slate
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, kd in enumerate(result.key_details):
            row = kd_table.add_row()
            bg_color = HEX_ALT_ROW_BG if row_idx % 2 == 1 else "FFFFFF"

            # Field
            c0 = row.cells[0]
            c0.width = kd_widths[0]
            set_cell_background(c0, bg_color)
            set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(kd.field)
            r0.bold = True
            r0.font.size = Pt(9)

            # Doc A
            c1 = row.cells[1]
            c1.width = kd_widths[1]
            set_cell_background(c1, bg_color)
            set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(kd.document_a_value)
            r1.font.size = Pt(9)

            # Doc B
            c2 = row.cells[2]
            c2.width = kd_widths[2]
            set_cell_background(c2, bg_color)
            set_cell_margins(c2, top=80, bottom=80, left=120, right=120)
            p2 = c2.paragraphs[0]
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(kd.document_b_value)
            r2.font.size = Pt(9)

            # Status
            c3 = row.cells[3]
            c3.width = kd_widths[3]
            set_cell_background(c3, bg_color)
            set_cell_margins(c3, top=80, bottom=80, left=120, right=120)
            p3 = c3.paragraphs[0]
            p3.paragraph_format.space_after = Pt(0)
            r3 = p3.add_run(kd.status.replace("_", " ").upper())
            r3.bold = True
            r3.font.size = Pt(8.5)
            r3.font.color.rgb = COLOR_SECONDARY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- 5. OVERALL CONCLUSION ---
    add_section_header("5. Overall Conclusion")
    p_conc = doc.add_paragraph()
    p_conc.paragraph_format.space_after = Pt(12)
    p_conc.paragraph_format.left_indent = Inches(0.2)
    r_conc = p_conc.add_run(result.overall_conclusion or "No conclusion generated.")
    r_conc.font.size = Pt(10.5)

    # --- 6. LIMITATIONS & UNCERTAINTIES ---
    add_section_header("6. Processing Limitations & Observed Uncertainties")
    if not result.limitations:
        doc.add_paragraph("No processing defects, unreadable pages, or textual ambiguities were reported.")
    else:
        for lim in result.limitations:
            p_lim = doc.add_paragraph(style="List Bullet")
            p_lim.paragraph_format.space_after = Pt(3)
            r_lim = p_lim.add_run(lim)
            r_lim.font.size = Pt(10)

    # --- 7. VERIFICATION NOTICE & DISCLAIMER ---
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    disc_p = doc.add_paragraph()
    disc_p.paragraph_format.space_before = Pt(14)
    disc_p.paragraph_format.space_after = Pt(4)
    disc_run = disc_p.add_run("Audit Verification Notice & Compliance Disclaimer:")
    disc_run.bold = True
    disc_run.font.size = Pt(9)
    disc_run.font.color.rgb = COLOR_MUTED

    notice_p = doc.add_paragraph()
    notice_p.paragraph_format.space_after = Pt(0)
    notice_run = notice_p.add_run(
        f"{APP_DISCLAIMER} All summaries, similarities, and differences in this report are "
        "AI-generated based strictly on extracted textual evidence. Users must review and verify all "
        "findings against the certified, original source documents prior to executing business, academic, "
        "or contractual decisions."
    )
    notice_run.font.size = Pt(8.5)
    notice_run.font.color.rgb = COLOR_MUTED
    notice_run.italic = True

    # Output to buffer and optionally file
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    if output_path:
        with open(output_path, "wb") as f:
            f.write(buffer.getvalue())

    return buffer
