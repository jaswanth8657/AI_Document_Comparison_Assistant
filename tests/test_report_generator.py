"""Automated test suite for report_generator.py.

Tests DOCX file creation, document structure, table formatting,
and graceful handling of empty or missing sections.
"""

import io
import os
import docx
import pytest

from report_generator import generate_docx_report
from schemas import (
    ComparisonResult,
    DifferenceItem,
    KeyDetailItem,
    SimilarityItem,
)


@pytest.fixture
def sample_comparison_result():
    """Provides a fully populated ComparisonResult instance."""
    return ComparisonResult(
        document_a_summary="Baseline remote work policy from 2024 requires 6 months tenure and $500 equipment stipend.",
        document_b_summary="Revised 2025 policy relaxes tenure to 3 months, increases stipend to $750, and requires hardware security keys.",
        similarities=[
            SimilarityItem(
                topic="VPN Usage",
                description="Both policies require company VPN connection.",
                evidence_a="connect through the secure company VPN",
                evidence_b="connect through the secure company VPN",
                pages_a=[2],
                pages_b=[2],
            )
        ],
        differences=[
            DifferenceItem(
                topic="Equipment Allowance",
                change_type="modified",
                document_a_value="$500.00",
                document_b_value="$750.00",
                impact="Additional $250 home office support.",
                evidence_a="stipend of $500.00",
                evidence_b="stipend of $750.00",
                pages_a=[2],
                pages_b=[2],
            ),
            DifferenceItem(
                topic="FIDO2 MFA Requirement",
                change_type="added",
                document_a_value="Not found in the provided text.",
                document_b_value="Mandatory FIDO2 hardware keys.",
                impact="Tightens endpoint security.",
                evidence_a="",
                evidence_b="FIDO2 hardware security keys",
                pages_a=[],
                pages_b=[2],
            )
        ],
        key_details=[
            KeyDetailItem(
                field="Equipment Stipend",
                document_a_value="$500.00",
                document_b_value="$750.00",
                status="changed",
                evidence_a="$500.00",
                evidence_b="$750.00",
                pages_a=[2],
                pages_b=[2],
            )
        ],
        overall_conclusion="The 2025 revision provides improved allowances and faster eligibility alongside upgraded cybersecurity standards.",
        limitations=["Scanned documents and OCR are not used here as both files contain digital text."],
        doc_a_name="policy_v1.pdf",
        doc_b_name="policy_v2.pdf",
        timestamp="2026-10-28 23:30:00",
        model_used="gemini-2.5-flash",
    )


def test_generate_docx_report_buffer(sample_comparison_result):
    """Verify that report generation produces a valid DOCX byte buffer."""
    buf = generate_docx_report(sample_comparison_result)
    assert isinstance(buf, io.BytesIO)
    raw_bytes = buf.getvalue()

    # Standard DOCX file begins with ZIP magic signature PK\x03\x04
    assert raw_bytes.startswith(b"PK\x03\x04")
    assert len(raw_bytes) > 5000


def test_docx_content_and_sections(sample_comparison_result):
    """Verify that all required sections, headings, and tables exist in generated DOCX."""
    buf = generate_docx_report(sample_comparison_result)
    doc = docx.Document(buf)

    # Collect all paragraph text
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "AI Document Comparison Assistant" in full_text
    assert "1. Executive Document Summaries" in full_text
    assert "2. Core Similarities & Shared Provisions" in full_text
    assert "3. Meaningful Differences & Modifications" in full_text
    assert "4. Key Details Comparison" in full_text
    assert "5. Overall Conclusion" in full_text
    assert "6. Processing Limitations & Observed Uncertainties" in full_text
    assert "Audit Verification Notice & Compliance Disclaimer" in full_text

    # Verify summary contents
    assert "Baseline remote work policy from 2024" in full_text
    assert "Revised 2025 policy relaxes tenure" in full_text

    # Verify tables exist (Metadata, Differences, Key Details)
    assert len(doc.tables) >= 3

    # Check differences table content
    table_texts = ["\n".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
    combined_table_text = "\n".join(table_texts)
    assert "Equipment Allowance" in combined_table_text
    assert "[MODIFIED]" in combined_table_text
    assert "$500.00" in combined_table_text
    assert "$750.00" in combined_table_text


def test_docx_empty_sections_handling():
    """Verify report generation handles empty similarities, differences, and limitations gracefully."""
    empty_result = ComparisonResult(
        document_a_summary="Summary A only.",
        document_b_summary="Summary B only.",
        similarities=[],
        differences=[],
        key_details=[],
        overall_conclusion="Documents are distinct.",
        limitations=[],
        doc_a_name="doc_a.pdf",
        doc_b_name="doc_b.pdf",
        timestamp="2026-10-28 23:30:00",
        model_used="gemini-2.5-flash",
    )

    buf = generate_docx_report(empty_result)
    doc = docx.Document(buf)
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "No significant common provisions" in full_text
    assert "No substantive differences" in full_text
    assert "No specific factual criteria" in full_text
    assert "No processing defects" in full_text


def test_docx_save_to_file(tmp_path, sample_comparison_result):
    """Verify that report can be written directly to a specified output file path."""
    dest_path = tmp_path / "test_report.docx"
    generate_docx_report(sample_comparison_result, output_path=str(dest_path))

    assert os.path.exists(dest_path)
    assert os.path.getsize(dest_path) > 5000

